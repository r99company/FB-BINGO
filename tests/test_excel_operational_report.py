from pathlib import Path

from openpyxl import load_workbook

from app.reports.excel_exporter import export_game_history


def test_export_game_history_creates_operational_workbook(tmp_path: Path):
    output = tmp_path / "reporte.xlsx"

    export_game_history(
        output,
        game_name="Rápida 1",
        series="0001",
        called_numbers=[7, 22, 45, 90],
        finished_at="2026-09-08 20:15:00",
        status="finalizada",
    )

    workbook = load_workbook(output, data_only=True)
    assert workbook.sheetnames == ["Partidas"]
    sheet = workbook["Partidas"]
    assert sheet.max_row == 2
    assert sheet["B2"].value == "Rápida 1"
    assert sheet["C2"].value == "0001"
    assert sheet["D2"].value == 4
    assert sheet["E2"].value == "7, 22, 45, 90"
    assert sheet["F2"].value == "finalizada"
